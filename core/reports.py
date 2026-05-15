import openpyxl
from openpyxl.styles import Font
from django.http import HttpResponse

def generate_excel_report(queryset, filename, headers, data_func):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte"

    # Add headers
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.font = Font(bold=True)

    # Add data
    for row_num, obj in enumerate(queryset, 2):
        data = data_func(obj)
        for col_num, value in enumerate(data, 1):
            ws.cell(row=row_num, column=col_num).value = str(value) if value is not None else ""

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    wb.save(response)
    return response
